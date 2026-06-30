"""
文件处理工具 - 支持多种文档格式的文件处理和验证
整合了文件验证功能，保持代码简洁和功能完整
"""

import hashlib
import tempfile
from pathlib import Path
from typing import Any, Dict, Optional, Tuple

import aiofiles
import magic
from fastapi import UploadFile

from src.core.logging import get_logger

logger = get_logger(__name__)


class FileProcessingError(Exception):
    """文件处理异常"""
    pass


class FileHandler:
    """文件处理器 - 整合了文件验证和处理功能"""

    # 支持的文件类型和MIME类型映射
    SUPPORTED_MIME_TYPES = {
        'text/plain': 'txt',
        'text/markdown': 'md',
        'text/x-markdown': 'md',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document': 'docx',
        'application/epub+zip': 'epub',
        # 图片类型
        'image/jpeg': 'image',
        'image/png': 'image',
        'image/gif': 'image',
        'image/webp': 'image',
        'image/bmp': 'image',
        # 视频类型
        'video/mp4': 'video',
        'video/quicktime': 'video',
    }

    # 文件扩展名映射
    SUPPORTED_EXTENSIONS = {
        '.txt': 'txt',
        '.md': 'md',
        '.markdown': 'md',
        '.docx': 'docx',
        '.epub': 'epub',
        # 图片扩展名
        '.jpg': 'image',
        '.jpeg': 'image',
        '.png': 'image',
        '.gif': 'image',
        '.webp': 'image',
        '.bmp': 'image',
        # 视频扩展名
        '.mp4': 'video',
        '.mov': 'video',
    }

    # 文件类型配置（整合自validators.py的有用配置）
    FILE_TYPE_CONFIG = {
        'txt': {
            'mime_types': ['text/plain'],
            'extensions': ['.txt'],
            'max_size': 50 * 1024 * 1024,  # 50MB
            'description': '纯文本文档'
        },
        'md': {
            'mime_types': ['text/markdown', 'text/plain'],
            'extensions': ['.md', '.markdown'],
            'max_size': 50 * 1024 * 1024,  # 50MB
            'description': 'Markdown文档'
        },
        'docx': {
            'mime_types': ['application/vnd.openxmlformats-officedocument.wordprocessingml.document'],
            'extensions': ['.docx'],
            'max_size': 100 * 1024 * 1024,  # 100MB
            'description': 'Word文档'
        },
        'epub': {
            'mime_types': ['application/epub+zip'],
            'extensions': ['.epub'],
            'max_size': 200 * 1024 * 1024,  # 200MB
            'description': 'EPUB电子书'
        },
        'image': {
            'mime_types': ['image/jpeg', 'image/png', 'image/gif', 'image/webp', 'image/bmp'],
            'extensions': ['.jpg', '.jpeg', '.png', '.gif', '.webp', '.bmp'],
            'max_size': 10 * 1024 * 1024,  # 10MB
            'description': '图片文件'
        },
        'video': {
            'mime_types': ['video/mp4', 'video/quicktime'],
            'extensions': ['.mp4', '.mov'],
            'max_size': 500 * 1024 * 1024,  # 500MB
            'description': '视频文件'
        }
    }

    # 危险文件扩展名黑名单（来自validators.py）
    DANGEROUS_EXTENSIONS = {
        '.exe', '.bat', '.cmd', '.com', '.pif', '.scr', '.vbs', '.js', '.jar',
        '.app', '.deb', '.pkg', '.dmg', '.rpm', '.msi', '.dll', '.so', '.dylib'
    }

    # 危险MIME类型黑名单（来自validators.py）
    DANGEROUS_MIME_TYPES = {
        'application/x-executable',
        'application/x-msdownload',
        'application/x-msdos-program',
        'application/x-shellscript',
        'application/javascript',
        'application/x-java-archive'
    }

    @classmethod
    def validate_filename(cls, filename: str) -> bool:
        """验证文件名是否安全有效（来自validators.py）"""
        if not filename:
            return False

        # 检查非法字符
        illegal_chars = ['<', '>', ':', '"', '|', '?', '*', '\x00']
        for char in illegal_chars:
            if char in filename:
                return False

        # 检查长度
        if len(filename) > 255:
            return False

        # 检查保留名称（Windows）
        reserved_names = ['CON', 'PRN', 'AUX', 'NUL'] + [f'COM{i}' for i in range(1, 10)] + [f'LPT{i}' for i in range(1, 10)]
        name_without_ext = Path(filename).stem.upper()
        if name_without_ext in reserved_names:
            return False

        return True

    @classmethod
    def validate_file_security(cls, filename: str, mime_type: Optional[str] = None) -> Optional[str]:
        """验证文件安全性（来自validators.py）"""
        file_ext = Path(filename).suffix.lower()

        # 检查危险扩展名
        if file_ext in cls.DANGEROUS_EXTENSIONS:
            return f"危险文件类型: {file_ext}"

        # 检查危险MIME类型
        if mime_type and mime_type in cls.DANGEROUS_MIME_TYPES:
            return f"危险文件类型: {mime_type}"

        return None

    @classmethod
    def get_file_type_from_extension(cls, filename: str) -> Optional[str]:
        if not filename:
            return None

        ext = Path(filename).suffix.lower()
        return cls.SUPPORTED_EXTENSIONS.get(ext)

    @classmethod
    def get_file_type_from_mime(cls, mime_type: str) -> Optional[str]:
        """从MIME类型获取文件类型 - 按照specification规范实现"""
        return cls.SUPPORTED_MIME_TYPES.get(mime_type)

    @classmethod
    async def validate_file(cls, file: UploadFile) -> Tuple[str, Dict[str, Any]]:
        """
        验证上传的文件 - 整合了安全检查和类型验证

        Args:
            file: 上传的文件

        Returns:
            Tuple[文件类型, 文件信息]

        Raises:
            FileProcessingError: 文件验证失败
        """
        if not file.filename:
            raise FileProcessingError("文件名不能为空")

        # 验证文件名安全性
        if not cls.validate_filename(file.filename):
            raise FileProcessingError("文件名包含非法字符或不符合规范")

        # 检查文件大小
        file.file.seek(0, 2)  # 移动到文件末尾
        file_size = file.file.tell()
        file.file.seek(0)  # 重置到文件开头

        if file_size == 0:
            raise FileProcessingError("文件不能为空")

        # 从扩展名推断文件类型
        file_type_ext = cls.get_file_type_from_extension(file.filename)
        if not file_type_ext:
            raise FileProcessingError(f"不支持的文件扩展名: {Path(file.filename).suffix}")

        # 检查文件类型大小限制
        file_config = cls.FILE_TYPE_CONFIG.get(file_type_ext)
        if file_config and file_size > file_config['max_size']:
            max_size_mb = file_config['max_size'] // (1024 * 1024)
            raise FileProcessingError(
                f"文件大小超过{file_config['description']}限制，最大允许 {max_size_mb}MB"
            )

        # 读取文件开头用于MIME类型检测
        file_content = file.file.read(1024)
        file.file.seek(0)  # 重置到文件开头

        # 检测MIME类型
        try:
            mime_type = magic.from_buffer(file_content, mime=True)
            file_type_mime = cls.get_file_type_from_mime(mime_type)
        except Exception as e:
            logger.warning(f"MIME类型检测失败: {e}")
            mime_type = None
            file_type_mime = None

        # 安全检查
        security_error = cls.validate_file_security(file.filename, mime_type)
        if security_error:
            raise FileProcessingError(security_error)

        # 验证文件类型一致性
        if file_type_mime and file_type_mime != file_type_ext:
            logger.warning(f"文件类型不匹配: 扩展名={file_type_ext}, MIME={file_type_mime or mime_type}")
            # 以MIME类型为准，如果支持的话
            if file_type_mime in cls.SUPPORTED_EXTENSIONS.values():
                file_type = file_type_mime
            else:
                raise FileProcessingError(f"不支持的文件类型: {mime_type}")
        else:
            file_type = file_type_ext

        # 计算文件哈希
        file_hash = await cls.calculate_file_hash(file)
        file.file.seek(0)  # 重置到文件开头

        file_info = {
            'filename': file.filename,
            'size': file_size,
            'content_type': file.content_type,
            'detected_mime': mime_type,
            'file_type': file_type,
            'file_hash': file_hash,
            'description': file_config['description'] if file_config else '',
        }

        return file_type, file_info

    @classmethod
    async def calculate_file_hash(cls, file: UploadFile) -> str:
        """计算文件的SHA-256哈希值"""
        file.file.seek(0)
        hash_sha256 = hashlib.sha256()

        # 分块读取文件以处理大文件
        chunk_size = 8192
        while True:
            chunk = file.file.read(chunk_size)
            if not chunk:
                break
            hash_sha256.update(chunk)

        file.file.seek(0)  # 重置到文件开头
        return hash_sha256.hexdigest()

    @classmethod
    async def save_temp_file(cls, file: UploadFile) -> Tuple[str, str]:
        """
        保存文件到临时目录

        Returns:
            Tuple[临时文件路径, 文件名]
        """
        # 创建临时文件
        suffix = Path(file.filename).suffix if file.filename else '.tmp'
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
            file_path = temp_file.name

        # 保存文件内容
        async with aiofiles.open(file_path, 'wb') as f:
            file.file.seek(0)
            while True:
                chunk = file.file.read(8192)
                if not chunk:
                    break
                await f.write(chunk)

        file.file.seek(0)  # 重置到文件开头
        return file_path, Path(file.filename).name


class TextFileReader:
    """纯粹的文本文件读取器 - 单一职责：只负责文件读取"""

    @staticmethod
    async def read_file(file_path: str, encoding: str = 'utf-8') -> str:
        """
        读取文本文件内容

        Args:
            file_path: 文件路径
            encoding: 文件编码

        Returns:
            文件内容

        Raises:
            FileProcessingError: 文件读取失败
        """
        try:
            async with aiofiles.open(file_path, 'r', encoding=encoding) as f:
                content = await f.read()
            return content
        except UnicodeDecodeError:
            # 尝试其他编码
            for alt_encoding in ['gbk', 'gb2312', 'latin-1']:
                try:
                    async with aiofiles.open(file_path, 'r', encoding=alt_encoding) as f:
                        content = await f.read()
                    logger.info(f"使用编码 {alt_encoding} 成功读取文件")
                    return content
                except UnicodeDecodeError:
                    continue

            raise FileProcessingError("无法解码文件内容，尝试了多种编码格式")


# 保持向后兼容的别名
TextFileHandler = TextFileReader


class MarkdownMetadataExtractor:
    """Markdown元数据提取器 - 单一职责：只负责提取Markdown元数据"""

    @staticmethod
    def extract_metadata(text: str) -> Dict[str, Any]:
        """
        提取Markdown元数据

        Args:
            text: Markdown文本内容

        Returns:
            元数据字典
        """
        import re

        # 提取标题
        titles = re.findall(r'^#+\s+(.+)$', text, re.MULTILINE)

        # 提取章节标题（# 和 ## 级别）
        chapter_titles = re.findall(r'^#{1,2}\s+(.+)$', text, re.MULTILINE)

        # 提取各级标题统计
        heading_stats = {}
        for i in range(1, 7):  # H1-H6
            heading_pattern = f'^{"#" * i}\\s+(.+)$'
            headings = re.findall(heading_pattern, text, re.MULTILINE)
            heading_stats[f'h{i}'] = len(headings)

        return {
            'titles': titles,
            'chapters': chapter_titles,
            'title_count': len(titles),
            'chapter_count': len(chapter_titles),
            'heading_stats': heading_stats,
        }


# 保持向后兼容的别名
MarkdownFileHandler = type('MarkdownFileHandler', (), {
    'read_markdown_file': lambda file_path: TextFileReader.read_file(file_path),
    'extract_metadata': MarkdownMetadataExtractor.extract_metadata,
})


class DocxReader:
    """Word文档读取器 - 单一职责：只负责读取DOCX文件"""

    @staticmethod
    async def read_file(file_path: str) -> str:
        """
        读取Word文档内容

        Args:
            file_path: 文件路径

        Returns:
            文档文本内容

        Raises:
            FileProcessingError: 读取失败
        """
        try:
            from docx import Document
            doc = Document(file_path)

            # 提取所有段落文本
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)

            # 提取表格内容
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text)
                    if row_text:
                        text_parts.append(' | '.join(row_text))

            return '\n\n'.join(text_parts)
        except ImportError:
            raise FileProcessingError("未安装python-docx库，无法处理Word文档")
        except Exception as e:
            raise FileProcessingError(f"读取Word文档失败: {str(e)}")


class DocxStructureExtractor:
    """Word文档结构提取器 - 单一职责：只负责提取DOCX结构"""

    @staticmethod
    def extract_structure(file_path: str) -> Dict[str, Any]:
        """
        提取Word文档结构

        Args:
            file_path: 文件路径

        Returns:
            文档结构信息
        """
        try:
            from docx import Document
            doc = Document(file_path)

            headings = []
            for paragraph in doc.paragraphs:
                if paragraph.style.name.startswith('Heading'):
                    level = paragraph.style.name.replace('Heading ', '')
                    headings.append({
                        'text': paragraph.text,
                        'level': int(level),
                    })

            # 统计各种元素
            paragraph_count = len([p for p in doc.paragraphs if p.text.strip()])
            table_count = len(doc.tables)
            heading_count = len(headings)

            return {
                'headings': headings,
                'heading_count': heading_count,
                'paragraph_count': paragraph_count,
                'table_count': table_count,
                'structure_valid': heading_count > 0,
            }
        except Exception as e:
            logger.error(f"提取Word文档结构失败: {e}")
            return {}


# 保持向后兼容的别名
DocxFileHandler = type('DocxFileHandler', (), {
    'read_docx_file': lambda file_path: DocxReader.read_file(file_path),
    'extract_structure': DocxStructureExtractor.extract_structure,
})


class EpubReader:
    """EPUB文件读取器 - 单一职责：只负责读取EPUB文件"""

    @staticmethod
    async def read_file(file_path: str) -> str:
        """
        读取EPUB文件内容

        Args:
            file_path: 文件路径

        Returns:
            电子书文本内容

        Raises:
            FileProcessingError: 读取失败
        """
        try:
            import ebooklib
            from ebooklib import epub
            from bs4 import BeautifulSoup

            book = epub.read_epub(file_path)
            text_parts = []

            # 提取所有章节内容
            for item in book.get_items():
                if item.get_type() == ebooklib.ITEM_DOCUMENT:
                    try:
                        soup = BeautifulSoup(item.get_content(), 'html.parser')
                        # 移除脚本和样式
                        for script in soup(["script", "style"]):
                            script.extract()
                        text = soup.get_text(separator='\n', strip=True)
                        if text.strip():
                            text_parts.append(text)
                    except Exception as e:
                        logger.warning(f"处理EPUB章节失败: {e}")
                        continue

            return '\n\n'.join(text_parts)
        except ImportError:
            raise FileProcessingError("未安装ebooklib和beautifulsoup4库，无法处理EPUB文件")
        except Exception as e:
            raise FileProcessingError(f"读取EPUB文件失败: {str(e)}")


class EpubMetadataExtractor:
    """EPUB元数据提取器 - 单一职责：只负责提取EPUB元数据"""

    @staticmethod
    def extract_metadata(file_path: str) -> Dict[str, Any]:
        """
        提取EPUB元数据

        Args:
            file_path: 文件路径

        Returns:
            电子书元数据
        """
        try:
            import ebooklib
            from ebooklib import epub

            book = epub.read_epub(file_path)
            metadata = {}

            # 获取基本元数据
            if book.get_metadata('DC', 'title'):
                metadata['title'] = book.get_metadata('DC', 'title')[0][0]
            if book.get_metadata('DC', 'creator'):
                metadata['creator'] = book.get_metadata('DC', 'creator')[0][0]
            if book.get_metadata('DC', 'language'):
                metadata['language'] = book.get_metadata('DC', 'language')[0][0]
            if book.get_metadata('DC', 'publisher'):
                metadata['publisher'] = book.get_metadata('DC', 'publisher')[0][0]
            if book.get_metadata('DC', 'date'):
                metadata['date'] = book.get_metadata('DC', 'date')[0][0]

            # 统计章节数
            chapters = [item for item in book.get_items()
                        if item.get_type() == ebooklib.ITEM_DOCUMENT]
            metadata['chapter_count'] = len(chapters)

            # 统计图片数量
            images = [item for item in book.get_items()
                      if item.get_type() == ebooklib.ITEM_IMAGE]
            metadata['image_count'] = len(images)

            # 验证文件完整性
            metadata['has_metadata'] = bool(metadata.get('title') or metadata.get('creator'))
            metadata['structure_valid'] = metadata['chapter_count'] > 0

            return metadata
        except Exception as e:
            logger.error(f"提取EPUB元数据失败: {e}")
            return {}


# 保持向后兼容的别名
EpubFileHandler = type('EpubFileHandler', (), {
    'read_epub_file': lambda file_path: EpubReader.read_file(file_path),
    'extract_metadata': EpubMetadataExtractor.extract_metadata,
})


# 文件处理器工厂
def get_file_handler(file_type: str):
    """获取对应的文件处理器 - 按照specification规范实现"""
    readers = {
        'txt': TextFileReader,
        'md': TextFileReader,  # Markdown文件也使用文本读取器
        'docx': DocxReader,
        'epub': EpubReader,
    }

    handler = readers.get(file_type)
    if not handler:
        raise FileProcessingError(f"不支持的文件类型: {file_type}")

    return handler


# 元数据提取器工厂
def get_metadata_extractor(file_type: str):
    """获取对应的元数据提取器"""
    extractors = {
        'txt': None,  # 纯文本文件没有特殊元数据
        'md': MarkdownMetadataExtractor,
        'docx': DocxStructureExtractor,
        'epub': EpubMetadataExtractor,
    }

    extractor = extractors.get(file_type)
    if not extractor:
        raise FileProcessingError(f"文件类型 {file_type} 不支持元数据提取")

    return extractor


__all__ = [
    # 核心类
    "FileHandler",
    "FileProcessingError",

    # 文件读取器 (单一职责)
    "TextFileReader",
    "DocxReader",
    "EpubReader",

    # 元数据提取器 (单一职责)
    "MarkdownMetadataExtractor",
    "DocxStructureExtractor",
    "EpubMetadataExtractor",

    # 工厂函数
    "get_file_handler",
    "get_metadata_extractor",

    # 向后兼容的别名
    "TextFileHandler",
    "MarkdownFileHandler",
    "DocxFileHandler",
    "EpubFileHandler",
]
